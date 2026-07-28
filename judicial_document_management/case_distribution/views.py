# case_distribution/views.py - ИСПРАВЛЕННАЯ ВЕРСИЯ

from rest_framework import viewsets, status, generics
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.contrib.contenttypes.models import ContentType
from django.db.models import Q
from django.utils import timezone
from django.shortcuts import get_object_or_404
import logging
from decimal import Decimal

from users.models import User
from .models import (
    DistributionSettings, JudgeSpecialization, DepartmentDistribution,
    JudgeLoad, CaseDistribution, DistributionLog, JudgeExclusion,
    SanctionJudge, ReassignmentLog, DistributionRule, ComplexitySetting
)
from .serializers import (
    DistributionSettingsSerializer, JudgeSpecializationSerializer,
    DepartmentDistributionSerializer, JudgeLoadSerializer,
    CaseDistributionSerializer, DistributionLogSerializer,
    JudgeExclusionSerializer, SanctionJudgeSerializer,
    ReassignmentLogSerializer, DistributionRuleSerializer,
    ComplexitySettingSerializer
)
from .services import DistributionService

logger = logging.getLogger(__name__)


class DistributionSettingsViewSet(viewsets.ModelViewSet):
    """Настройки распределения"""
    queryset = DistributionSettings.objects.all()
    serializer_class = DistributionSettingsSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        return self.queryset.all()
    
    def perform_create(self, serializer):
        if DistributionSettings.objects.exists():
            from rest_framework import serializers
            raise serializers.ValidationError("Настройки уже существуют")
        serializer.save()


class JudgeSpecializationViewSet(viewsets.ModelViewSet):
    """Специализации судей"""
    queryset = JudgeSpecialization.objects.all()
    serializer_class = JudgeSpecializationSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['user', 'specialization', 'is_active']
    search_fields = ['user__first_name', 'user__last_name', 'user__username']


class DepartmentDistributionViewSet(viewsets.ModelViewSet):
    """Коллегии судей"""
    queryset = DepartmentDistribution.objects.all()
    serializer_class = DepartmentDistributionSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['department_type', 'is_active']
    search_fields = ['name', 'description']


class JudgeLoadViewSet(viewsets.ReadOnlyModelViewSet):
    """Нагрузка судей"""
    queryset = JudgeLoad.objects.all()
    serializer_class = JudgeLoadSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['judge', 'year', 'month']
    
    @action(detail=False, methods=['get'])
    def current(self, request):
        now = timezone.now()
        service = DistributionService(user=request.user)
        stats = service.get_statistics(now.year, now.month)
        return Response(stats)


class CaseDistributionViewSet(viewsets.ModelViewSet):
    queryset = CaseDistribution.objects.all()
    serializer_class = CaseDistributionSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['status', 'source', 'assigned_judge', 'case_category']
    search_fields = ['case_number']

    # ========== КРИТИЧЕСКОЕ ИСПРАВЛЕНИЕ ==========
    # Действия, для которых НЕ нужно фильтровать по show_all
    DETAIL_ACTIONS = {
        'retrieve', 'update', 'partial_update', 'destroy',
        'set_complexity', 'update_complexity', 'view_case', 'reassign'
    }

    def get_queryset(self):
        queryset = super().get_queryset()
        
        # ========== ИСПРАВЛЕНИЕ: Для действий с конкретным объектом ==========
        # НЕ применяем фильтр show_all - нужно найти объект по ID
        if self.action in self.DETAIL_ACTIONS:
            return queryset
        
        # ========== Для списковых действий применяем фильтры ==========
        show_all = self.request.query_params.get('show_all', 'false').lower() == 'true'
        
        # Фильтр по статусу
        status_param = self.request.query_params.get('status')
        if status_param:
            if status_param == 'pending':
                queryset = queryset.filter(assigned_judge__isnull=True)
            elif status_param == 'distributed':
                queryset = queryset.filter(assigned_judge__isnull=False)
        
        # По умолчанию показываем только дела без судьи
        if not show_all:
            queryset = queryset.filter(assigned_judge__isnull=True)
        
        # Поиск по номеру дела
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(case_number__icontains=search)
        
        # Фильтр по типу дела
        case_type = self.request.query_params.get('case_type')
        if case_type:
            queryset = queryset.filter(case_category=case_type)
        
        # Фильтр по судье
        judge_id = self.request.query_params.get('judge_id')
        if judge_id:
            queryset = queryset.filter(assigned_judge_id=judge_id)
        
        # Фильтр по сложности
        complexity_min = self.request.query_params.get('complexity_min')
        if complexity_min:
            try:
                queryset = queryset.filter(complexity__gte=Decimal(complexity_min))
            except (ValueError, TypeError):
                pass
        
        complexity_max = self.request.query_params.get('complexity_max')
        if complexity_max:
            try:
                queryset = queryset.filter(complexity__lte=Decimal(complexity_max))
            except (ValueError, TypeError):
                pass
        
        # Фильтр по датам
        date_from = self.request.query_params.get('date_from')
        if date_from:
            queryset = queryset.filter(distributed_at__date__gte=date_from)
        
        date_to = self.request.query_params.get('date_to')
        if date_to:
            queryset = queryset.filter(distributed_at__date__lte=date_to)
        
        return queryset

    # ========== УСТАНОВКА СЛОЖНОСТИ ==========
    @action(
        detail=True,
        methods=['post', 'patch'],
        url_path='set-complexity',
        url_name='set-complexity'
    )
    def set_complexity(self, request, pk=None):
        """
        Установить сложность конкретного дела.

        POST/PATCH: /case-distribution/cases/{id}/set-complexity/
        Body: { "complexity": 1 }
        """
        distribution = self.get_object()

        complexity_value = request.data.get('complexity')

        if complexity_value is None:
            return Response(
                {
                    'error': 'Не указана сложность',
                    'detail': 'Параметр complexity обязателен'
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            complexity = Decimal(str(complexity_value))
        except (ValueError, TypeError, ArithmeticError):
            return Response(
                {'error': 'Некорректное значение сложности'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if complexity < Decimal('0.5') or complexity > Decimal('10'):
            return Response(
                {'error': 'Сложность должна быть от 0.5 до 10'},
                status=status.HTTP_400_BAD_REQUEST
            )

        old_complexity = distribution.complexity or Decimal('1')

        distribution.complexity = complexity
        distribution.save(update_fields=['complexity', 'updated_at'])

        # Обновляем нагрузку судьи
        if distribution.assigned_judge:
            try:
                now = timezone.now()
                load = JudgeLoad.objects.filter(
                    judge=distribution.assigned_judge,
                    year=now.year,
                    month=now.month
                ).first()

                if load:
                    diff = complexity - Decimal(str(old_complexity))
                    load.total_complexity = Decimal(str(load.total_complexity or 0)) + diff
                    load.save(update_fields=['total_complexity', 'last_updated'])

                    logger.info(
                        f"Обновлена нагрузка судьи {distribution.assigned_judge}: "
                        f"сложность {old_complexity} -> {complexity}"
                    )
            except Exception:
                logger.exception(
                    'Ошибка обновления нагрузки судьи при изменении сложности дела %s',
                    distribution.pk
                )

        return Response(
            {
                'success': True,
                'message': f'Сложность дела обновлена на {complexity}',
                'distribution': self.get_serializer(distribution).data
            },
            status=status.HTTP_200_OK
        )

    # ========== ОБНОВЛЕНИЕ ДЕЛА ЧЕРЕЗ СТАНДАРТНЫЙ PATCH ==========
    def update(self, request, *args, **kwargs):
        """
        Переопределяем update для поддержки обновления сложности через стандартный PATCH.
        URL: /case-distribution/cases/{pk}/
        """
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        
        # Проверяем, есть ли complexity в данных
        complexity = request.data.get('complexity')
        if complexity is not None:
            try:
                new_complexity = Decimal(str(complexity))
                if Decimal('0.5') <= new_complexity <= Decimal('10'):
                    old_complexity = instance.complexity or Decimal('1')
                    instance.complexity = new_complexity
                    
                    # Обновляем нагрузку судьи
                    if instance.assigned_judge:
                        now = timezone.now()
                        load = JudgeLoad.objects.filter(
                            judge=instance.assigned_judge,
                            year=now.year,
                            month=now.month
                        ).first()
                        if load:
                            diff = new_complexity - old_complexity
                            load.total_complexity = Decimal(str(load.total_complexity or 0)) + diff
                            load.save(update_fields=['total_complexity', 'last_updated'])
                            
                            logger.info(
                                f"Обновлена нагрузка судьи {instance.assigned_judge}: "
                                f"сложность {old_complexity} -> {new_complexity}"
                            )
                else:
                    return Response(
                        {'error': 'Сложность должна быть от 0.5 до 10'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
            except (ValueError, TypeError):
                return Response(
                    {'error': 'Некорректное значение сложности'},
                    status=status.HTTP_400_BAD_REQUEST
                )
        
        # Продолжаем стандартное обновление
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        
        return Response(serializer.data)

    @action(detail=True, methods=['get'], url_path='view')
    def view_case(self, request, pk=None):
        """Получить данные дела для просмотра."""
        distribution = self.get_object()
        serializer = self.get_serializer(distribution)
        return Response(serializer.data)

    @action(detail=False, methods=['post'])
    def auto_distribute(self, request):
        """Автоматическое распределение дел"""
        service = DistributionService(user=request.user)
        case_ids = request.data.get('case_ids', [])
        
        if not isinstance(case_ids, list):
            case_ids = []
        
        if not case_ids:
            return Response(
                {'error': 'Не указаны дела для распределения. Передайте case_ids.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        distributions = CaseDistribution.objects.filter(
            id__in=case_ids,
            assigned_judge__isnull=True
        ).select_related('case_content_type')
        
        if not distributions.exists():
            return Response(
                {'error': 'Нет дел для распределения (все уже распределены или неверные ID)'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        cases = []
        for dist in distributions:
            if dist.case:
                cases.append(dist.case)
        
        if not cases:
            return Response(
                {'error': 'Не удалось получить объекты дел'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        log = DistributionLog.objects.create(
            executed_by=request.user,
            distribution_type='auto'
        )
        
        result = service.distribute_cases(cases, log)
        
        return Response({
            'log': DistributionLogSerializer(log).data,
            'result': result
        })

    @action(detail=False, methods=['post'])
    def manual_distribute(self, request):
        """Ручное распределение дела"""
        case_id = request.data.get('case_id')
        judge_id = request.data.get('judge_id')
        
        if not case_id or not judge_id:
            return Response(
                {'error': 'Не указаны дело и судья'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        try:
            judge = User.objects.get(id=judge_id, role='judge')
        except User.DoesNotExist:
            return Response(
                {'error': 'Судья не найден'},
                status=status.HTTP_404_NOT_FOUND
            )
        
        try:
            distribution = CaseDistribution.objects.get(id=case_id)
        except CaseDistribution.DoesNotExist:
            return Response(
                {'error': 'Распределение не найдено'},
                status=status.HTTP_404_NOT_FOUND
            )
        
        if distribution.assigned_judge:
            return Response(
                {'error': 'Дело уже распределено'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        if distribution.case:
            service = DistributionService(user=request.user)
            service._set_presiding_judge(distribution.case, judge)
        
        log = DistributionLog.objects.create(
            executed_by=request.user,
            distribution_type='manual'
        )
        
        distribution.assign_judge(
            judge=judge,
            user=request.user,
            source='manual',
            log_entry=log
        )
        
        log.total_cases = 1
        log.distributed_cases = 1
        log.mark_completed()
        
        return Response({
            'distribution': CaseDistributionSerializer(distribution).data,
            'log': DistributionLogSerializer(log).data
        })
    
    @action(detail=True, methods=['post'])
    def reassign(self, request, pk=None):
        """Перераспределение дела другому судье"""
        distribution = self.get_object()
        new_judge_id = request.data.get('new_judge_id')
        reason = request.data.get('reason')
        reason_type = request.data.get('reason_type', 'other')
        
        if not new_judge_id:
            return Response(
                {'error': 'Не указан новый судья'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        if not reason:
            return Response(
                {'error': 'Не указана причина перераспределения'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        try:
            new_judge = User.objects.get(id=new_judge_id, role='judge')
        except User.DoesNotExist:
            return Response(
                {'error': 'Судья не найден'},
                status=status.HTTP_404_NOT_FOUND
            )
        
        service = DistributionService(user=request.user)
        result = service.reassign_case(distribution, new_judge, reason, reason_type)
        
        if 'error' in result:
            return Response(
                {'error': result['error']},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        return Response(result)

    @action(detail=False, methods=['get'], url_path='available_judges')
    def available_judges(self, request):
        """Получить список доступных судей"""
        service = DistributionService(user=request.user)
        date = request.query_params.get('date')
        case_id = request.query_params.get('case_id')
        case_type = request.query_params.get('case_type')
        
        case = None
        if case_id:
            case = self._get_case_by_id(case_id, case_type or 'criminal')
        
        available = service.get_available_judges(
            date=date,
            case=case,
            exclude_judge=request.query_params.get('exclude_judge')
        )
        
        result = []
        for item in available:
            judge = item['judge']
            load_info = item['load_info']
            result.append({
                'id': judge.id,
                'username': judge.username,
                'first_name': judge.first_name,
                'last_name': judge.last_name,
                'full_name': judge.get_full_name(),
                'current_load': float(load_info.get('current_complexity', 0)),
                'total_load': float(load_info.get('total_complexity', 0)),
                'load_percentage': load_info.get('load_percentage', 100),
                'max_cases': load_info.get('max_cases', 100),
                'effective_max': load_info.get('effective_max_cases', 100),
            })
        
        return Response(result)

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        """Получить статистику распределения"""
        service = DistributionService(user=request.user)
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        
        if year:
            year = int(year)
        else:
            year = timezone.now().year
            
        if month:
            month = int(month)
        else:
            month = timezone.now().month
        
        stats = service.get_statistics(year, month)
        
        pending = CaseDistribution.objects.filter(assigned_judge__isnull=True).count()
        distributed = CaseDistribution.objects.filter(assigned_judge__isnull=False).count()
        
        active_judges = User.objects.filter(role='judge', is_active=True)
        service_local = DistributionService()
        available = service_local.get_available_judges()
        active_count = len(available)
        
        by_case_type = {}
        case_categories = CaseDistribution._meta.get_field('case_category').choices
        
        for category, label in case_categories:
            count = CaseDistribution.objects.filter(
                case_category=category,
                assigned_judge__isnull=False
            ).count()
            if count > 0:
                by_case_type[label] = count
        
        return Response({
            'total_cases': CaseDistribution.objects.count(),
            'without_judge': pending,
            'distributed': distributed,
            'active_judges': active_count,
            'by_case_type': by_case_type,
            'monthly': stats
        })

    @action(detail=False, methods=['get'], url_path='pending')
    def pending_cases(self, request):
        """Получить дела, ожидающие распределения (без судьи)"""
        queryset = CaseDistribution.objects.filter(assigned_judge__isnull=True)
        
        search = request.query_params.get('search')
        if search:
            queryset = queryset.filter(case_number__icontains=search)
        
        case_type = request.query_params.get('case_type')
        if case_type:
            queryset = queryset.filter(case_category=case_type)
        
        ordering = request.query_params.get('ordering', '-created_at')
        queryset = queryset.order_by(ordering)
        
        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

    def _get_case_by_id(self, case_id, case_type='criminal'):
        """Получить объект дела по ID и типу"""
        from criminal_proceedings.models import CriminalProceedings
        from civil_proceedings.models import CivilProceedings
        from administrative_code.models import AdministrativeProceedings
        from administrative_proceedings.models import KasProceedings
        
        model_map = {
            'criminal': CriminalProceedings,
            'civil': CivilProceedings,
            'coap': AdministrativeProceedings,
            'kas': KasProceedings,
        }
        
        model = model_map.get(case_type)
        if not model:
            return None
        
        try:
            return model.objects.get(id=case_id)
        except model.DoesNotExist:
            return None


class DistributionLogViewSet(viewsets.ReadOnlyModelViewSet):
    """Журнал распределения"""
    queryset = DistributionLog.objects.all()
    serializer_class = DistributionLogSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['status', 'distribution_type', 'executed_by']
    search_fields = ['protocol_number']
    
    @action(detail=True, methods=['get'])
    def export(self, request, pk=None):
        log = self.get_object()
        return self._export_to_word(log)
    
    @action(detail=False, methods=['get'], url_path='export')
    def export_all(self, request):
        from io import BytesIO
        from docx import Document
        from django.http import FileResponse
        from datetime import datetime
        
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        action_type = request.query_params.get('action_type')
        
        queryset = self.get_queryset()
        
        if start_date:
            try:
                start = datetime.strptime(start_date, '%Y-%m-%d')
                queryset = queryset.filter(started_at__date__gte=start)
            except ValueError:
                pass
        
        if end_date:
            try:
                end = datetime.strptime(end_date, '%Y-%m-%d')
                queryset = queryset.filter(started_at__date__lte=end)
            except ValueError:
                pass
        
        if action_type:
            queryset = queryset.filter(distribution_type=action_type)
        
        queryset = queryset.order_by('-started_at')
        
        if not queryset.exists():
            return Response(
                {'error': 'Нет записей для экспорта'},
                status=status.HTTP_404_NOT_FOUND
            )
        
        doc = Document()
        
        title = doc.add_heading('ЖУРНАЛ РАСПРЕДЕЛЕНИЯ ДЕЛ', 0)
        title.alignment = 1
        
        period_info = []
        if start_date:
            period_info.append(f"с {start_date}")
        if end_date:
            period_info.append(f"по {end_date}")
        
        if period_info:
            doc.add_paragraph(f"Период: {' '.join(period_info)}")
        
        doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        doc.add_paragraph(f"Всего записей: {queryset.count()}")
        
        doc.add_heading('Протоколы распределения', level=1)
        
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        
        header_cells = table.rows[0].cells
        header_cells[0].text = '№ протокола'
        header_cells[1].text = 'Дата/время'
        header_cells[2].text = 'Кто выполнил'
        header_cells[3].text = 'Тип'
        header_cells[4].text = 'Статус'
        header_cells[5].text = 'Всего дел'
        header_cells[6].text = 'Распределено'
        
        for cell in header_cells:
            cell.paragraphs[0].runs[0].bold = True
        
        for log in queryset:
            row_cells = table.add_row().cells
            row_cells[0].text = log.protocol_number or '-'
            row_cells[1].text = log.started_at.strftime('%d.%m.%Y %H:%M') if log.started_at else '-'
            row_cells[2].text = log.executed_by.get_full_name() if log.executed_by else '-'
            row_cells[3].text = log.get_distribution_type_display()
            row_cells[4].text = log.get_status_display()
            row_cells[5].text = str(log.total_cases or 0)
            row_cells[6].text = str(log.distributed_cases or 0)
        
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        filename = f"Журнал_распределения_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
        
        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename
        )
    
    def _export_to_word(self, log):
        from io import BytesIO
        from docx import Document
        from django.http import FileResponse
        
        doc = Document()
        
        title = doc.add_heading('ПРОТОКОЛ РАСПРЕДЕЛЕНИЯ ДЕЛ', 0)
        title.alignment = 1
        
        doc.add_paragraph(f'Протокол №: {log.protocol_number}')
        doc.add_paragraph(f'Дата и время: {log.started_at.strftime("%d.%m.%Y %H:%M:%S")}')
        doc.add_paragraph(f'Выполнил: {log.executed_by.get_full_name() if log.executed_by else "—"}')
        doc.add_paragraph(f'Тип: {log.get_distribution_type_display()}')
        
        doc.add_heading('Статистика', level=1)
        
        stats = [
            ('Всего дел', log.total_cases),
            ('Распределено', log.distributed_cases),
            ('Не удалось распределить', log.failed_cases),
            ('Перераспределено', log.reassigned_cases),
        ]
        
        for label, value in stats:
            doc.add_paragraph(f'{label}: {value}')
        
        if log.judges_info:
            doc.add_heading('Нагрузка судей', level=1)
            
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            
            header_cells = table.rows[0].cells
            header_cells[0].text = 'Судья'
            header_cells[1].text = 'Количество дел'
            
            for judge_name, count in log.judges_info.items():
                row_cells = table.add_row().cells
                row_cells[0].text = judge_name
                row_cells[1].text = str(count)
        
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return FileResponse(
            buffer,
            as_attachment=True,
            filename=f'protokol_{log.protocol_number}.docx'
        )


class JudgeExclusionViewSet(viewsets.ModelViewSet):
    queryset = JudgeExclusion.objects.all()
    serializer_class = JudgeExclusionSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['judge', 'exclusion_type', 'is_permanent']


class SanctionJudgeViewSet(viewsets.ModelViewSet):
    queryset = SanctionJudge.objects.all()
    serializer_class = SanctionJudgeSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['sanction_type', 'is_active']


class DistributionRuleViewSet(viewsets.ModelViewSet):
    queryset = DistributionRule.objects.all()
    serializer_class = DistributionRuleSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['rule_type', 'category', 'is_active']
    search_fields = ['name', 'description']


class ComplexitySettingViewSet(viewsets.ModelViewSet):
    queryset = ComplexitySetting.objects.all()
    serializer_class = ComplexitySettingSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['category', 'is_active']


class ReassignmentLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = ReassignmentLog.objects.all()
    serializer_class = ReassignmentLogSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['from_judge', 'to_judge', 'reason_type']
    search_fields = ['case_distribution__case_number']
